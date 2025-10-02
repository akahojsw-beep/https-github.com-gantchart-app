import streamlit as st
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
import calendar
import io

st.title("ExcelからガントチャートPowerPoint生成")

uploaded_file = st.file_uploader("Excelファイルをアップロードしてください", type=["xlsx"])

if uploaded_file is not None:
    # Excelデータ読み込み
    df = pd.read_excel(uploaded_file)
    for col in ["Start1", "End1", "Start2", "End2", "Start3", "End3", "Start4", "End4"]:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col])

    # 年月リスト作成
    start_dates = pd.concat([df.get("Start1", pd.Series([])), df.get("Start2", pd.Series([])), df.get("Start3", pd.Series([])), df.get("Start4", pd.Series([]))]).dropna()
    end_dates = pd.concat([df.get("End1", pd.Series([])), df.get("End2", pd.Series([])), df.get("End3", pd.Series([])), df.get("End4", pd.Series([]))]).dropna()
    start_ym = start_dates.min().to_period("M")-2
    end_ym = end_dates.max().to_period("M") + 1
    ym_list = pd.period_range(start_ym, end_ym, freq='M')
    years = [str(ym.year) for ym in ym_list]
    months = [str(ym.month) for ym in ym_list]

    # pptxテンプレート
    prs = Presentation('template.pptx')
    slide = prs.slides[0]
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    margin_x = Inches(0.2)
    margin_y_top = Inches(1.5)
    margin_y_bottom = Inches(1)

    cell_widths = [Inches(1.0), Inches(0.7), Inches(0.7), Inches(1.6)]
    year_label_height = Inches(0.5)
    month_label_height = Inches(0.4)

    chart_left = margin_x + sum(cell_widths)
    chart_top = margin_y_top + year_label_height + month_label_height
    num_sites = len(df)
    num_months = len(ym_list)
    chart_width = slide_width - chart_left - margin_x
    chart_height = slide_height - chart_top - margin_y_bottom
    gantt_cell_width = chart_width / num_months
    row_height = chart_height / num_sites

    # 列名ラベル
    header_labels = ['パタン', 'スペック', 'サイト']
    header_top = margin_y_top
    header_height = year_label_height + month_label_height

    for i, label in enumerate(header_labels):
        left = margin_x + sum(cell_widths[:i+1])
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, header_top, cell_widths[i+1], header_height
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(230, 230, 230)
        rect.line.color.rgb = RGBColor(180, 180, 180)
        rect.line.width = Pt(1)
        tf = rect.text_frame
        tf.text = label
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        tf.paragraphs[0].font.name = "Yu Gothic UI"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # 年ラベル
    prev_year = ""
    year_start_idx = 0
    for idx, year in enumerate(years):
        if year != prev_year:
            if prev_year != "":
                left = chart_left + gantt_cell_width * year_start_idx
                width = gantt_cell_width * (idx - year_start_idx)
                rect = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, left, margin_y_top, width, year_label_height
                )
                rect.fill.solid()
                rect.fill.fore_color.rgb = RGBColor(230, 230, 230)
                rect.line.color.rgb = RGBColor(180, 180, 180)
                rect.line.width = Pt(1)
                tf = rect.text_frame
                tf.text = prev_year
                tf.paragraphs[0].font.size = Pt(10)
                tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
                tf.paragraphs[0].font.name = "Yu Gothic UI"
                tf.paragraphs[0].font.bold = True
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            prev_year = year
            year_start_idx = idx
    left = chart_left + gantt_cell_width * year_start_idx
    width = gantt_cell_width * (num_months - year_start_idx)
    rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, margin_y_top, width, year_label_height
    )
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(230, 230, 230)
    rect.line.color.rgb = RGBColor(180, 180, 180)
    rect.line.width = Pt(1)
    tf = rect.text_frame
    tf.text = prev_year
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
    tf.paragraphs[0].font.name = "Yu Gothic UI"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # 月ラベル
    for i, month in enumerate(months):
        left = chart_left + gantt_cell_width * i
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, margin_y_top + year_label_height, gantt_cell_width, month_label_height
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(230, 230, 230)
        rect.line.color.rgb = RGBColor(180, 180, 180)
        rect.line.width = Pt(1.0)
        tf = rect.text_frame
        tf.text = month
        tf.paragraphs[0].font.size = Pt(9.0)
        tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        tf.paragraphs[0].font.name = "Yu Gothic UI"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # セル結合風の表示
    def make_merge_groups(df, col_name):
        merge_info = []
        prev = None
        start_idx = 0
        for idx, val in enumerate(df[col_name]):
            if prev is None:
                prev = val
                start_idx = idx
            elif val != prev:
                merge_info.append((start_idx, idx-1, prev))
                prev = val
                start_idx = idx
        merge_info.append((start_idx, len(df)-1, prev))
        return merge_info

    size_groups = make_merge_groups(df, 'サイズ')
    patan_groups = make_merge_groups(df, 'パタン')
    spec_groups = make_merge_groups(df, 'スペック')

    # サイズ
    for start, end, val in size_groups:
        top = chart_top + row_height * start
        height = row_height * (end - start + 1)
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, margin_x, top, cell_widths[0], height
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(230, 230, 230)
        rect.line.color.rgb = RGBColor(180, 180, 180)
        rect.line.width = Pt(1)
        tf = rect.text_frame
        tf.text = str(val)
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        tf.paragraphs[0].font.name = "Yu Gothic UI"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # パタン
    for start, end, val in patan_groups:
        top = chart_top + row_height * start
        height = row_height * (end - start + 1)
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, margin_x + cell_widths[0], top, cell_widths[1], height
        )
        rect.fill.background()
        rect.line.color.rgb = RGBColor(180, 180, 180)
        rect.line.width = Pt(1)
        tf = rect.text_frame
        tf.text = str(val)
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        tf.paragraphs[0].font.name = "Yu Gothic UI"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # スペック
    for start, end, val in spec_groups:
        top = chart_top + row_height * start
        height = row_height * (end - start + 1)
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, margin_x + cell_widths[0] + cell_widths[1], top, cell_widths[2], height
        )
        rect.fill.background()
        rect.line.color.rgb = RGBColor(180, 180, 180)
        rect.line.width = Pt(1)
        tf = rect.text_frame
        tf.text = str(val)
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        tf.paragraphs[0].font.name = "Yu Gothic UI"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # サイト名
    site_pos_dict = {}
    for idx, row in df.iterrows():
        top = chart_top + row_height * idx
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, margin_x + cell_widths[0] + cell_widths[1] + cell_widths[2], top, cell_widths[3], row_height
        )
        rect.fill.background()
        rect.line.color.rgb = RGBColor(180, 180, 180)
        rect.line.width = Pt(1)
        tf = rect.text_frame
        tf.text = str(row['サイト'])
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        tf.paragraphs[0].font.name = "Yu Gothic UI"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        site_pos_dict[(row['サイズ'], row['パタン'], row['スペック'], row['サイト'])] = idx

    # 縦線・横線
    for i in range(num_months + 1):
        x = chart_left + gantt_cell_width * i
        is_solid = (i == 0) or (i == num_months)
        if i > 0 and i < num_months:
            prev_month = int(months[i-1])
            curr_month = int(months[i])
            if prev_month == 12 and curr_month == 1:
                is_solid = True
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, int(x), int(chart_top), int(x), int(chart_top + chart_height)
        )
        line.line.color.rgb = RGBColor(180, 180, 180)
        line.line.width = Pt(0.75)
        line.line.transparency = 0
        if is_solid:
            line.line.dash_style = MSO_LINE_DASH_STYLE.SOLID
        else:
            line.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    for i in range(num_sites + 1):
        y = chart_top + row_height * i
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, int(chart_left), int(y), int(chart_left + chart_width), int(y)
        )
        line.line.color.rgb = RGBColor(180, 180, 180)
        line.line.width = Pt(1.5)
        line.line.transparency = 0
        line.line.dash_style = MSO_LINE_DASH_STYLE.SOLID

    # ガントバー
    bar_labels = ["生産", "輸送・装着", "走行", "量産準備期間"]
    bar_short = ["S1", "S2", "S3", "S4"]
    bar_colors = [
        RGBColor(184, 204, 228),
        RGBColor(255, 242, 204),
        RGBColor(163, 201, 236),
        RGBColor(184, 204, 228)
    ]
    for idx, row in df.iterrows():
        for j, (start_col, end_col) in enumerate([
            ("Start1", "End1"),
            ("Start2", "End2"),
            ("Start3", "End3"),
            ("Start4", "End4")
        ]):
            if (start_col in df.columns) and (end_col in df.columns) and pd.notnull(row[start_col]) and pd.notnull(row[end_col]):
                start = row[start_col]
                end = row[end_col]
                start_ym = start.to_period("M")
                end_ym = end.to_period("M")
                start_idx = ym_list.get_loc(start_ym)
                end_idx = ym_list.get_loc(end_ym)
                days_in_start_month = calendar.monthrange(start.year, start.month)[1]
                days_in_end_month = calendar.monthrange(end.year, end.month)[1]
                start_offset = (start.day - 1) / days_in_start_month
                end_offset = end.day / days_in_end_month

                bar_left = chart_left + gantt_cell_width * (start_idx + start_offset)
                bar_right = chart_left + gantt_cell_width * end_idx + gantt_cell_width * end_offset
                bar_width = bar_right - bar_left

                site_row = idx
                bar_height = row_height * 0.27
                bar_top = chart_top + row_height * site_row + (row_height - bar_height) / 2

                shape = slide.shapes.add_shape(
                    MSO_SHAPE.PENTAGON, bar_left, bar_top, bar_width, bar_height
                )
                fill = shape.fill
                fill.solid()
                fill.fore_color.rgb = bar_colors[j]
                shape.line.color.rgb = RGBColor(180, 180, 180)
                shape.line.width = Pt(1)
                shape.line.transparency = 0

                tf = shape.text_frame
                tf.text = bar_labels[j]
                tf.paragraphs[0].font.size = Pt(8.5)
                tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
                tf.paragraphs[0].font.name = "Yu Gothic UI"
                tf.paragraphs[0].font.bold = True
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER

 # S1, S2, S3, S4の値を取得してガントバー下に表示
            s_col = bar_short[j]
            months_label = ""
            if s_col in row and pd.notnull(row[s_col]):
                months_label = f"{row[s_col]:.1f}ヶ月"
            label_top = bar_top + bar_height + Pt(1)
            label_height = Pt(12)
            label_shape = slide.shapes.add_textbox(
                bar_left, label_top, bar_width, label_height
            )
            label_tf = label_shape.text_frame
            label_tf.text = months_label
            label_tf.paragraphs[0].font.size = Pt(8)
            label_tf.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
            label_tf.paragraphs[0].font.name = "Yu Gothic UI"
            label_tf.paragraphs[0].font.bold = True
            label_tf.paragraphs[0].alignment = PP_ALIGN.CENTER



    # pptx保存（BytesIOで一時保存）
    output = io.BytesIO()
    prs.save(output)
    output.seek(0)

    st.success("PowerPointファイルが生成されました")
    st.download_button(
        label="PowerPointファイルをダウンロード",
        data=output,
        file_name="gantt_chart.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )